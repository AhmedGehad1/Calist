"""Check that every form's fields land correctly — across the whole archive.

    python tools/audit.py "D:/MedCal Pro"                  run the audit
    python tools/audit.py --diff BEFORE.json.gz AFTER.json.gz

**Read-only against the forms.** It opens every inspection form under the
``Customers <year>`` folders, and writes nothing anywhere except its own report
folder, ``<root>/_Calist audit``. The customer forms are the signed record; the
audit must never be the thing that changes one.

Why this exists
---------------
Every misread this project has had was *silent*. A form re-laid-out between
visits reads a plausible wrong value, or a blank, and nothing flags it — a blank
field looks exactly like one an engineer left empty. Unit tests cannot find
these: they need the real forms. So this reads every form the way the register
does, then asks the form, independently, what it says:

1. **The Certificate tab.** Nearly every form carries one, with an
   ``Equipment Data`` block whose values are *formulas naming the device-tab
   cell they copy* — ``Serial No. = Data!K15``. When rows are inserted, Excel
   rewrites those references, so each certificate is a per-form statement of
   where the fields really are. Where it names a different cell than Calist's
   map and the values differ, the map is wrong on that form.
2. **The device tab's own printed labels**, via ``calist.locate_by_labels``.
3. **Shape**: a date is never a serial, a caption is never a model.
4. **Cross-year**: the same serial on the same device should read the same
   maker and model every year it is inspected.
5. **Calibrator contamination**: the certificates also name the instrument that
   did the calibrating; a device field holding one of those is reading the
   wrong block.

Each disagreement is classified as a MAP ERROR (Calist reads the wrong cell —
fixable here), a FORM ERROR (the form itself is wrong or empty — a person has
to correct it), EXPECTED (explained, not a fault), or UNDETERMINED.

Speed
-----
A process pool, which the app deliberately does not use: the app reads one
form in ~2 ms and a pool costs more than it saves, but this reads every form
two extra ways (~100 ms each), so across 80,000 forms the pool is what makes it
minutes rather than hours.
"""

from __future__ import annotations

import argparse
import gzip
import json
import os
import random
import re
import sys
import time
import warnings
from collections import Counter, defaultdict
from datetime import datetime
from multiprocessing import Pool
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
if str(REPO) not in sys.path:
    sys.path.insert(0, str(REPO))

warnings.filterwarnings("ignore")

import calist                                       # noqa: E402
from calist import (                                # noqa: E402
    _CELL_RE, _LABEL_COLUMNS, _LABEL_MAX_ROW, _NOT_A_VERDICT, _PLACEHOLDERS,
    _is_a_label, _open_workbook, _read_with, classify_serial, clean,
    locate_by_labels, normalise_status, read_best,
)
from device_config import DEVICE_CONFIGS            # noqa: E402

try:
    from device_names import DEVICE_NAMES           # noqa: E402
except ImportError:                                 # pragma: no cover
    DEVICE_NAMES = {}

YEARS = ("2023", "2024", "2025", "2026")

#: (device, year) pairs left out of the audit by the owner's decision —
#: Digital blood pressure 2023 is a known set of problem forms, not to fix.
SKIPPED = {("CB", "2023")}
REPORT_DIR = "_Calist audit"

#: Every field the register carries that a form can get wrong.
CHECKED = ("Manufacturer", "Model", "S.N", "S.N2", "Location", "Date",
           "Status", "Status2")

#: The subset the certificate's Equipment Data block also states.
CERT_FIELDS = ("Manufacturer", "Model", "S.N", "Location", "Date")

#: The certificate tab. "Certificate", "ECG.certificate", and on FC just "cert".
_CERT_TAB = re.compile(r"cert", re.I)

_EQUIPMENT_DATA = re.compile(r"^\s*equipment\s*data\s*$", re.I)
_CERT_LABELS = {
    "S.N": re.compile(r"^serial\s*(no\.?|number)?\s*:?\s*$|^s\.?\s*n\.?\s*:?\s*$", re.I),
    "Model": re.compile(r"^model\s*:?\s*$", re.I),
    "Manufacturer": re.compile(r"^manufacturer\s*:?\s*$", re.I),
    "Location": re.compile(r"^location\s*:?\s*$", re.I),
    "Type": re.compile(r"^equipment\s*type\s*:?\s*$", re.I),
}
#: The certificate's calibration date sits in the "Calibration Data" block,
#: above Equipment Data, so it is searched for everywhere.
_CERT_DATE = re.compile(r"^calib\.?\s*date\s*:?\s*$", re.I)
_CALIBRATOR_ROW = re.compile(r"^calibration\s*device\s*:?\s*$", re.I)

#: The first cross-sheet reference in a formula: Data!K15, 'Device data'!$K$15.
_FORMULA_REF = re.compile(
    r"(?:'((?:[^']|'')+)'|([A-Za-z0-9_. ]+))!\$?([A-Z]{1,3})\$?(\d+)")

#: Real month names only. A looser "three to nine letters" read the models
#: BSM-3562 and Erkameter3000 as dates.
_MONTH = (r"(jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?")
_DATE_TEXT = re.compile(
    r"^\s*(\d{1,2}\s*[-/.]\s*\d{1,2}\s*[-/.]\s*\d{2,4}"
    r"|\d{4}\s*[-/.]\s*\d{1,2}\s*[-/.]\s*\d{1,2}(\s+00:00:00)?"
    rf"|\d{{1,2}}[\s-]*{_MONTH}[\s,-]*\d{{2,4}}"
    rf"|{_MONTH}[\s-]*\d{{1,2}}[\s,-]+\d{{2,4}})\s*$", re.I)
#: A bare reading in a Location or Model box ("90.0", "25.8"). Wider than the
#: reader's own rule on purpose: this only judges, it never rejects a layout.
_MEASUREMENT = re.compile(r"^\s*\d{1,3}\.\d{1,2}\s*$")

#: How xlrd hands back a date typed into a real date cell: the bare serial.
_EXCEL_SERIAL = re.compile(r"^\s*[34]\d{4}(\.0+)?\s*$")

GRID = {f"{c}{r}" for r in range(1, _LABEL_MAX_ROW + 1) for c in _LABEL_COLUMNS}
CERT_GRID = {f"{c}{r}" for r in range(1, 81) for c in _LABEL_COLUMNS}


# ──────────────────────────────────────────────────────────────────────────────
# Reading one form every way (runs in a worker process)
# ──────────────────────────────────────────────────────────────────────────────

def _map_used(config: dict, how: str) -> dict[str, str] | None:
    """The cells read_best actually read, when it was a written-down layout."""
    if how == "primary":
        return config["cells"]
    match = re.match(r"alt (\d+)$", how)
    if match:
        alternates = config.get("alt_cells", [])
        index = int(match.group(1)) - 1
        if 0 <= index < len(alternates):
            return {**config["cells"], **alternates[index]}
    return None


def _value_right_of(values: dict[str, str], ref: str, label: str) -> str | None:
    """The first cell right of a label that is neither empty, a repeat of the
    label across its own merged span, nor another caption."""
    match = re.match(r"([A-Z]+)(\d+)$", ref)
    if not match or match.group(1) not in _LABEL_COLUMNS:
        return None
    column, row = match.group(1), match.group(2)
    for other in _LABEL_COLUMNS[_LABEL_COLUMNS.index(column) + 1:]:
        candidate = values.get(f"{other}{row}", "")
        if not candidate or candidate == label:
            continue
        if _is_a_label(candidate):
            return None
        return f"{other}{row}"
    return None


def _formula_of(source, ref: str) -> tuple[str, str | None, str | None]:
    """How a certificate cell got its value: typed, or a formula — and if a
    formula, which sheet and cell it copies (None when it cannot be told)."""
    data = getattr(source, "_data", None)
    if data is None:
        return "unknown", None, None                # .xls: formulas not visible
    anchor = source._anchor(ref)
    at = data.find(b'<c r="' + anchor.encode() + b'"')
    match = _CELL_RE.match(data, at) if at >= 0 else None
    body = (match.group(2) or b"") if match else b""
    if b"<f" not in body:
        return "typed", None, None
    text = re.search(rb"<f[^>]*>(.*?)</f>", body, re.S)
    if not text:
        return "formula", None, None                # shared formula, no text
    found = _FORMULA_REF.search(calist._unescape(text.group(1)))
    if not found:
        return "formula", None, None
    sheet = (found.group(1) or found.group(2) or "").replace("''", "'").strip()
    return "formula", sheet, f"{found.group(3)}{found.group(4)}"


def _read_certificate(source) -> dict | None:
    """The certificate's Equipment Data block, and the calibrator it names."""
    tab = next((t for t in getattr(source, "sheet_names", [])
                if _CERT_TAB.search(t)), None)
    if tab is None or not source.select_sheet(tab):
        return None
    values = {ref: clean(value) for ref, value in source.values(CERT_GRID).items()}
    rows = [int(re.sub(r"\D", "", ref)) for ref, text in values.items()
            if text and _EQUIPMENT_DATA.match(text)]
    block = min(rows) if rows else None

    fields: dict[str, list] = {}
    calibrator: list[str] = []
    for ref in sorted(values, key=lambda r: (int(re.sub(r"\D", "", r)), len(r), r)):
        text = values[ref]
        if not text:
            continue
        row = int(re.sub(r"\D", "", ref))
        if "Date" not in fields and _CERT_DATE.match(text):
            at = _value_right_of(values, ref, text)
            if at:
                how, sheet, cell = _formula_of(source, at)
                fields["Date"] = [values[at], at, how, sheet, cell]
            continue
        if _CALIBRATOR_ROW.match(text):
            # The instrument that did the calibrating: "Model :" and "S.N :"
            # further along the same row. Harvested so that a device field
            # holding one of these can be recognised as a misread.
            for other in _LABEL_COLUMNS:
                cell = values.get(f"{other}{row}", "")
                if re.match(r"^(model|s\.?\s*n\.?)\s*:?\s*$", cell, re.I):
                    at = _value_right_of(values, f"{other}{row}", cell)
                    if at:
                        calibrator.append(values[at])
            continue
        if block is None or row < block:
            continue                                # only the device's own block
        for field, pattern in _CERT_LABELS.items():
            if field in fields or not pattern.match(text):
                continue
            at = _value_right_of(values, ref, text)
            if at:
                how, sheet, cell = _formula_of(source, at)
                fields[field] = [values[at], at, how, sheet, cell]
            break
    return {"tab": tab, "block": block is not None, "fields": fields,
            "calibrator": calibrator}


def audit_form(job: tuple[str, str, str]) -> dict:
    """Everything the audit needs to know about one form."""
    path, year, code = job
    config = DEVICE_CONFIGS[code]
    out: dict = {"p": path, "y": year, "c": code}
    try:
        record, how = read_best(path, config)
    except calist.NotAForm as error:
        out["notform"] = str(error)
        return out
    except Exception as error:                      # noqa: BLE001
        out["err"] = f"{type(error).__name__}: {error}"[:200]
        return out
    out["how"] = how
    out["land"] = {field: clean(record.get(field)) for field in CHECKED
                   if field in record}

    try:
        source = _open_workbook(path)
    except Exception as error:                      # noqa: BLE001
        out["err2"] = f"{type(error).__name__}: {error}"[:200]
        return out
    try:
        device_tab = getattr(source, "sheet_name", "")
        out["tab"] = device_tab
        used = _map_used(config, how)
        anchor = getattr(source, "_anchor", None)
        if used and anchor:
            out["mapA"] = {f: anchor(r) for f, r in used.items() if r}

        located = locate_by_labels(source.values(GRID))
        if located:
            read = _read_with(source, located)
            out["lab"] = {f: [read.get(f, ""), r] for f, r in located.items()}

        certificate = _read_certificate(source)
        if certificate:
            # Resolve formula targets that point back at the device tab to the
            # anchor of their merged box, so "L17" and "K17" of one box agree.
            targets = {f: v[4] for f, v in certificate["fields"].items()
                       if v[3] and v[3].lower() == device_tab.lower() and v[4]}
            if targets and anchor and source.select_sheet(device_tab):
                for field, cell in targets.items():
                    certificate["fields"][field].append(anchor(cell))
            out["cert"] = certificate
    except Exception as error:                      # noqa: BLE001
        out["err2"] = f"{type(error).__name__}: {error}"[:200]
    finally:
        source.close()
    return out


# ──────────────────────────────────────────────────────────────────────────────
# Judging
# ──────────────────────────────────────────────────────────────────────────────

def _norm(value: str) -> str:
    text = " ".join((value or "").strip().lower().split())
    return re.sub(r"\.0+$", "", text)               # 2085.0 is 2085


def _date_key(value: str):
    """A date as (y, m, d) where it can be read as one, else the text.

    A test period ("30/04/2023 - 18/06/2023", which a certificate prints as
    the form does) keys as its start, which is what the register shows.
    """
    text = (value or "").strip()
    period = calist._PERIOD_RE.match(text)
    if period:
        text = text[:period.end(3)]
    match = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", text)
    if match:
        return tuple(int(g) for g in match.groups())
    match = re.match(r"^(\d{1,2})\s*[-/.]\s*(\d{1,2})\s*[-/.]\s*(\d{2,4})$", text)
    if match:
        d, m, y = (int(g) for g in match.groups())
        return (y + 2000 if y < 100 else y, m, d)
    return _norm(text)


def same(field: str, a: str, b: str) -> bool:
    if field == "Date":
        return _date_key(a) == _date_key(b)
    return _norm(a) == _norm(b)


def is_date(value: str, field: str = "Date") -> bool:
    """A date as these forms write it.

    A leading apostrophe is an engineer forcing text ("'10-02-2024") and
    still a date. The bare number xlrd returns for a real .xls date cell
    ("45306.0") counts only in the Date column: in Model or Serial a five-digit
    number is a number — "46036.0" on a Centrifuge form is not a date.
    """
    text = (value or "").strip().lstrip("'")
    if calist.is_calendar_date(text) or _DATE_TEXT.match(text) and not             calist._DATE_LIKE_RE.match(text):
        return True
    return field == "Date" and bool(_EXCEL_SERIAL.match(text)
                                    or calist._DATE_TIME_RE.match(text))


def shape(field: str, value: str) -> str:
    """'good', 'blank', 'placeholder', or a reason the value cannot be right."""
    text = (value or "").strip()
    if not text:
        return "blank"
    low = text.lower()
    if field in ("Status", "Status2"):
        if normalise_status(text):
            return "good"
        if text in ("0", "0.0"):
            # The verdict box copies an empty test sheet: empty, not wrong.
            return "blank"
        if _is_a_label(text):
            return "a caption, not a verdict"
        if low.rstrip(".") in _NOT_A_VERDICT:
            return "answers a different question"
        return "not a recognised verdict"
    if field in ("S.N", "S.N2"):
        kind = classify_serial(text)
        if kind in ("real", "assigned"):
            return "good"
        if kind == "placeholder":
            return "placeholder"
        if is_date(text, field):
            return "a date, not a serial"
        if _is_a_label(text):
            return "a caption, not a serial"
        if normalise_status(text):
            return "a verdict, not a serial"
        return "no digits — not a serial"
    if low in _PLACEHOLDERS:
        return "placeholder"
    if _is_a_label(text):
        return "a caption"
    if field == "Date":
        return "good" if is_date(text, "Date") else "not a date"
    if field in ("Model", "Manufacturer") and is_date(text, field):
        return "a date"
    if field == "Manufacturer" and re.fullmatch(r"[\d.\-/ ]+", text):
        return "a number, not a maker"
    if field in ("Location", "Model") and _MEASUREMENT.match(text):
        return "a measurement, not a " + field.lower()
    return "good"


def judge(form: dict) -> dict[str, list]:
    """Per field: [verdict, reason, better value or '', evidence]."""
    code = form["c"]
    config = DEVICE_CONFIGS[code]
    land = form.get("land", {})
    lab = form.get("lab", {})
    cert = (form.get("cert") or {}).get("fields", {})
    map_anchor = form.get("mapA", {})
    wanted = [f for f in CHECKED if config["cells"].get(f)]
    out: dict[str, list] = {}
    model, serial = land.get("Model", ""), land.get("S.N", "")
    same_twice = bool(model) and _norm(model) not in _PLACEHOLDERS         and _norm(model) == _norm(serial)

    for field in wanted:
        value = land.get(field, "")
        kind = shape(field, value)
        if same_twice and field in ("Model", "S.N"):
            kind = "the same text as the " + ("serial" if field == "Model" else "model")
        good = kind == "good"

        c = cert.get(field)
        c_val = c[0] if c else ""
        c_good = bool(c) and shape(field, c_val) == "good"
        c_how = c[2] if c else None
        c_target = (c[5] if c and len(c) > 5 else None)
        c_tab = c[3] if c else None
        l = lab.get(field)
        l_val = l[0] if l else ""
        l_good = bool(l) and shape(field, l_val) == "good"
        m_ref = map_anchor.get(field)

        # BB-style second serials: the certificate names the probe. The
        # register shows device serial first, probe in brackets — as asked.
        if (field == "S.N" and c and c_target and config["cells"].get("S.N2")
                and map_anchor.get("S.N2") == c_target):
            out[field] = ["EXPECTED", "certificate names the second serial",
                          "", f"cert -> {c_target}"]
            continue

        if good:
            if c and same(field, value, c_val):
                out[field] = ["OK", "certificate agrees", "", ""]
            elif l and same(field, value, l_val):
                out[field] = ["OK", "form's label agrees", "", ""]
            elif form.get("how") == "none" and field not in ("Status", "Status2"):
                # No known layout fits this form, so read_best fell back to the
                # primary map's reading of a sheet it does not describe. Text
                # there passes every shape check — "2. Protective Earth
                # Resistance" is a perfectly ordinary string — and is still not
                # the manufacturer. Unconfirmed means untrusted here.
                out[field] = ["WRONG VALUE", "no known layout fits this form",
                              c_val if c_good else (l_val if l_good else ""),
                              m_ref or ""]
            elif c and c_good and c_how == "formula" and c_target and m_ref \
                    and c_target != m_ref:
                out[field] = ["MAP ERROR",
                              f"certificate copies {c_target}, map reads {m_ref}",
                              c_val, f"{m_ref}->{c_target}"]
            elif c and c_good and c_how == "formula" and c_tab and c_tab.lower() \
                    != (form.get("tab") or "").lower():
                out[field] = ["MAP ERROR",
                              f"certificate copies '{c_tab}'!{c[4]}, not the tab read",
                              c_val, f"{m_ref}->'{c_tab}'!{c[4]}"]
            elif c and c_how == "typed" and not c_good:
                out[field] = ["OK", "certificate holds a typed placeholder", "", ""]
            elif c and c_good:
                out[field] = ["UNDETERMINED",
                              "certificate says something else", c_val,
                              f"cert {c_how}"]
            elif l and l_good:
                out[field] = ["UNDETERMINED", "form's label says something else",
                              l_val, f"label {l[1]}"]
            elif (l and _norm(l_val) in _PLACEHOLDERS) or (
                    c and _norm(c_val) in _PLACEHOLDERS and c_how != "typed"):
                # The form's own label (or a certificate copying it) says N.A
                # and Calist has something else: a Baby Warmer read its ward
                # as the Model exactly like this. Doubtful, not fine.
                out[field] = ["UNDETERMINED", "the form's label says N.A",
                              "N.A", f"label {l[1]}" if l else "cert"]
            else:
                out[field] = ["OK", "shape only", "", ""]
            continue

        # Not a good value. Can the form tell us what it should have been?
        if c and c_good and not (c_how == "formula" and c_target and m_ref
                                 and c_target == m_ref):
            where = (f"{m_ref}->{c_target}" if c_target
                     else f"{m_ref}->'{c_tab}'!{c[4]}" if c_tab else "cert")
            out[field] = ["MAP ERROR", f"Calist has {kind}; the form has a value",
                          c_val, where]
        elif l and l_good and (not m_ref or l[1] != m_ref):
            out[field] = ["MAP ERROR", f"Calist has {kind}; the form's label has a value",
                          l_val, f"{m_ref}->{l[1]}"]
        elif kind in ("blank", "placeholder"):
            # Nothing anywhere on the form: the engineer left it empty or N.A.
            # A gap in the record, not a misread — kept apart so the real
            # errors are not buried under it.
            out[field] = ["MISSING", f"{kind} on the form", "", ""]
        else:
            # Cannot be right, and nothing says where the right value is.
            # One such form is an engineer's slip; the same thing across a
            # device's forms is its map — the report tells the two apart.
            out[field] = ["WRONG VALUE", f"the cell holds {kind}", "", m_ref or ""]
    return out


#: A device-year where at least this share of forms have a WRONG VALUE in one
#: field is reported as a map problem rather than as engineers' slips.
SYSTEMIC_SHARE = 0.2


# ──────────────────────────────────────────────────────────────────────────────
# The run
# ──────────────────────────────────────────────────────────────────────────────

def discover(root: str):
    known, unknown = [], []
    for year in YEARS:
        folder = os.path.join(root, f"Customers {year}")
        if not os.path.isdir(folder):
            continue
        for path in calist.find_source_files(folder):
            outcome = calist.classify_file(path)
            if outcome.status == calist.READY:
                if (outcome.device_code, year) in SKIPPED:
                    continue
                known.append((path, year, outcome.device_code))
            else:
                unknown.append((path, year, outcome.device_code, outcome.status))
    return known, unknown


def unknown_kind(path: str, code: str | None) -> str:
    stem = Path(path).stem
    low = stem.lower()
    if re.search(r"device\s*_?\s*list|^list\b|^devicelist|فشلت|قائمة", low) \
            or code in ("DEVICE", "LIST", "DEVICELIST"):
        return "not a form (a list or register)"
    if re.search(r"(^|[-\s])[a-z]{1,5}0{2,4}(\b|[-\s]|$)", low) or "datrend" in low:
        return "blank template (unit 000)"
    if code and code in DEVICE_NAMES and calist._CODES_RE.match(stem):
        return "real device code, not mapped"
    return "form named without a device code"


def run(root: str, workers: int, limit: int | None,
        out: str | None = None) -> Path:
    started = time.monotonic()
    known, unknown = discover(root)
    if limit:
        random.Random(1).shuffle(known)
        known = known[:limit]
    print(f"forms to audit: {len(known):,}   unknown-code files: {len(unknown):,}",
          flush=True)

    results = []
    with Pool(workers) as pool:
        for index, form in enumerate(pool.imap_unordered(audit_form, known,
                                                         chunksize=16), 1):
            results.append(form)
            if index % 2000 == 0 or index == len(known):
                rate = index / (time.monotonic() - started)
                left = (len(known) - index) / rate if rate else 0
                print(f"  {index:>6,} / {len(known):,}   {rate:5.0f}/s   "
                      f"~{left/60:4.1f} min left", flush=True)

    for form in results:
        if "land" in form:
            form["judge"] = judge(form)

    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    out_dir = Path(out) if out else Path(root) / REPORT_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    snapshot = out_dir / f"audit-{stamp}.json.gz"
    for year in YEARS:
        forms = (Path(root) / f"Customers {year}").resolve()
        if forms in snapshot.resolve().parents:
            raise RuntimeError(f"refusing to write the snapshot inside {forms}")
    with gzip.open(snapshot, "wt", encoding="utf-8") as handle:
        json.dump({"root": root, "at": stamp, "forms": results,
                   "unknown": unknown}, handle, ensure_ascii=False)
    workbook = out_dir / f"audit-{stamp}.xlsx"
    write_report(results, unknown, workbook, root)
    print(f"\nsnapshot  {snapshot}\nworkbook  {workbook}\n"
          f"took      {(time.monotonic() - started) / 60:.1f} min", flush=True)
    return workbook


# ──────────────────────────────────────────────────────────────────────────────
# The workbook
# ──────────────────────────────────────────────────────────────────────────────

def _link(cell, path: str) -> None:
    cell.value = os.path.basename(path)
    cell.hyperlink = "file:///" + path.replace("\\", "/")
    cell.style = "Hyperlink"


def _sheet(workbook, title: str, header: list[str], widths: list[int]):
    from openpyxl.styles import Font, PatternFill
    sheet = workbook.create_sheet(title)
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    sheet.freeze_panes = "A2"
    return sheet


def cross_year(results: list[dict]) -> list[tuple]:
    """Same device code, same real serial, different maker or model."""
    seen = defaultdict(list)
    for form in results:
        land = form.get("land") or {}
        serial = land.get("S.N", "")
        if classify_serial(serial) != "real" or len(serial) < 5:
            continue
        seen[(form["c"], _norm(serial))].append(form)
    conflicts = []
    for (code, serial), forms in seen.items():
        if len({f["y"] for f in forms}) < 2:
            continue
        for field in ("Manufacturer", "Model"):
            readings = {_norm(f["land"].get(field, "")) for f in forms
                        if shape(field, f["land"].get(field, "")) == "good"}
            if len(readings) > 1:
                conflicts.append((code, serial, field, forms))
    return conflicts


def calibrator_hits(results: list[dict]) -> list[tuple]:
    vocabulary = Counter()
    for form in results:
        for value in (form.get("cert") or {}).get("calibrator", []):
            if len(_norm(value)) >= 3 and _norm(value) not in ("n.a", "na"):
                vocabulary[_norm(value)] += 1
    common = {v for v, n in vocabulary.items() if n >= 5}
    hits = []
    for form in results:
        for field in ("Manufacturer", "Model", "S.N"):
            value = (form.get("land") or {}).get(field, "")
            if value and _norm(value) in common:
                hits.append((form, field, value, vocabulary[_norm(value)]))
    return hits


def write_report(results: list[dict], unknown: list[tuple], path: Path,
                 root: str) -> None:
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    workbook.remove(workbook.active)
    judged = [f for f in results if "judge" in f]
    not_forms = [f for f in results if "notform" in f]
    unreadable = [f for f in results if "judge" not in f and "notform" not in f]

    # ── Summary: device × field verdicts ──────────────────────────────────
    verdicts = ("OK", "MAP ERROR", "WRONG VALUE", "MISSING", "UNDETERMINED",
                "EXPECTED")
    summary = _sheet(workbook, "Summary",
                     ["Code", "Device", "Forms", "Field", *verdicts, "% OK"],
                     [8, 30, 8, 12, 8, 11, 12, 10, 13, 10, 8])
    per = defaultdict(Counter)
    forms_per = Counter()
    forms_per_year = Counter()
    for form in judged:
        forms_per[form["c"]] += 1
        forms_per_year[(form["c"], form["y"])] += 1
        for field, verdict in form["judge"].items():
            per[(form["c"], field)][verdict[0]] += 1
    grand = Counter()
    for (code, field), counts in sorted(per.items()):
        total = sum(counts.values())
        grand.update(counts)
        summary.append([code, DEVICE_CONFIGS[code]["device_name"], forms_per[code],
                        field, *(counts[v] for v in verdicts),
                        round(counts["OK"] / total * 100, 1)])
    summary.auto_filter.ref = summary.dimensions

    # A WRONG VALUE that recurs across a device-year is its map, not a slip.
    wrong = defaultdict(list)
    for form in judged:
        for field, verdict in form["judge"].items():
            if verdict[0] == "WRONG VALUE":
                wrong[(form["c"], form["y"], field)].append(form)
    systemic = {key for key, forms in wrong.items()
                if len(forms) >= 3
                and len(forms) / forms_per_year[key[:2]] >= SYSTEMIC_SHARE}

    # ── Map errors, grouped into one proposed fix each ─────────────────────
    groups = defaultdict(list)
    for form in judged:
        for field, (verdict, reason, better, where) in form["judge"].items():
            if verdict == "MAP ERROR":
                groups[(form["c"], field, where, "evidence")].append((form, better))
            elif verdict == "WRONG VALUE" and (form["c"], form["y"], field) in systemic:
                groups[(form["c"], field, f"{where}: {reason}", "systemic")].append(
                    (form, ""))
    mapped = _sheet(workbook, "Map errors",
                    ["Code", "Device", "Field", "Map reads -> form says",
                     "Basis", "Forms", "% of this device", "Years",
                     "Calist wrote", "Form says", "Example"],
                    [8, 26, 12, 34, 10, 7, 10, 20, 26, 26, 40])
    for (code, field, where, basis), items in sorted(groups.items(),
                                                    key=lambda kv: -len(kv[1])):
        form, better = items[0]
        row = [code, DEVICE_CONFIGS[code]["device_name"], field, where, basis,
               len(items), round(len(items) / forms_per[code] * 100, 1),
               ", ".join(sorted({f["y"] for f, _ in items})),
               form["land"].get(field, ""), better, ""]
        mapped.append(row)
        _link(mapped.cell(mapped.max_row, 11), form["p"])
    mapped.auto_filter.ref = mapped.dimensions

    # ── One row per problem, for the forms themselves ─────────────────────
    for title, wanted in (("Form errors", "WRONG VALUE"),
                          ("Missing on form", "MISSING"),
                          ("Undetermined", "UNDETERMINED")):
        sheet = _sheet(workbook, title,
                       ["Year", "Code", "Device", "Field", "Calist wrote", "Why",
                        "Form elsewhere says", "Evidence", "File", "Folder"],
                       [6, 7, 24, 11, 26, 34, 26, 16, 34, 60])
        for form in judged:
            for field, (verdict, reason, better, where) in form["judge"].items():
                if verdict != wanted:
                    continue
                if wanted == "WRONG VALUE" and (form["c"], form["y"], field) in systemic:
                    continue                        # reported as a map problem
                sheet.append([form["y"], form["c"],
                              DEVICE_CONFIGS[form["c"]]["device_name"], field,
                              form["land"].get(field, ""), reason, better, where,
                              "", os.path.dirname(form["p"]).replace(root, "")])
                _link(sheet.cell(sheet.max_row, 9), form["p"])
        sheet.auto_filter.ref = sheet.dimensions

    # ── Cross-year conflicts ───────────────────────────────────────────────
    sheet = _sheet(workbook, "Cross-year conflicts",
                   ["Code", "Serial", "Field", "Year", "Calist wrote", "File"],
                   [8, 22, 13, 6, 30, 40])
    for code, serial, field, forms in cross_year(judged):
        for form in sorted(forms, key=lambda f: f["y"]):
            sheet.append([code, serial, field, form["y"],
                          form["land"].get(field, ""), ""])
            _link(sheet.cell(sheet.max_row, 6), form["p"])
    sheet.auto_filter.ref = sheet.dimensions

    # ── Calibrator contamination ───────────────────────────────────────────
    sheet = _sheet(workbook, "Calibrator in device field",
                   ["Year", "Code", "Field", "Calist wrote",
                    "Seen as a calibrator on N forms", "File"],
                   [6, 7, 13, 26, 16, 40])
    for form, field, value, count in calibrator_hits(judged):
        sheet.append([form["y"], form["c"], field, value, count, ""])
        _link(sheet.cell(sheet.max_row, 6), form["p"])

    # ── Unknown codes ──────────────────────────────────────────────────────
    sheet = _sheet(workbook, "Unknown codes",
                   ["Kind", "Code", "Forms", "2023", "2024", "2025", "2026",
                    "Master-list name", "Example"],
                   [34, 18, 7, 6, 6, 6, 6, 30, 50])
    tally = defaultdict(lambda: [Counter(), None])
    # Never `path` here: that name is this function's report destination, and
    # reusing it once made the report save over the last customer file listed.
    for file, year, code, _status in unknown:
        key = (unknown_kind(file, code), code or "(none)")
        tally[key][0][year] += 1
        tally[key][1] = tally[key][1] or file
    for (kind, code), (years, example) in sorted(
            tally.items(), key=lambda kv: (kv[0][0], -sum(kv[1][0].values()))):
        sheet.append([kind, code, sum(years.values()), years["2023"],
                      years["2024"], years["2025"], years["2026"],
                      DEVICE_NAMES.get(code, ""), ""])
        _link(sheet.cell(sheet.max_row, 9), example)
    sheet.auto_filter.ref = sheet.dimensions

    # ── Device lists and other workbooks that are not forms ────────────────
    sheet = _sheet(workbook, "Not forms", ["Year", "Code", "Why", "File"],
                   [6, 7, 50, 50])
    for form in not_forms:
        sheet.append([form["y"], form["c"], form["notform"], ""])
        _link(sheet.cell(sheet.max_row, 4), form["p"])

    # ── Unreadable ─────────────────────────────────────────────────────────
    sheet = _sheet(workbook, "Unreadable", ["Year", "Code", "Error", "File"],
                   [6, 7, 60, 40])
    for form in unreadable:
        sheet.append([form["y"], form["c"], form.get("err", ""), ""])
        _link(sheet.cell(sheet.max_row, 4), form["p"])

    # ── Sample for a human to confirm: 3 per device per year ───────────────
    header = ["Year", "Code", "Device", "File", "Route"]
    for field in ("Manufacturer", "Model", "S.N", "Status", "Location", "Date",
                  "S.N2", "Status2"):
        header += [f"{field}", f"{field} (certificate)"]
    header += ["Audit verdicts", "Confirm", "Your notes"]
    sample = _sheet(workbook, "Sample to confirm", header,
                    [6, 7, 20, 34, 9] + [18, 18] * 8 + [40, 10, 40])
    rng = random.Random(7)
    strata = defaultdict(list)
    for form in judged:
        strata[(form["c"], form["y"])].append(form)
    picked = []
    for key in sorted(strata):
        forms = strata[key]
        picked += rng.sample(forms, min(3, len(forms)))
    confirm = DataValidation(type="list", formula1='"OK,Wrong"', allow_blank=True)
    sample.add_data_validation(confirm)
    for form in picked:
        cert = (form.get("cert") or {}).get("fields", {})
        row = [form["y"], form["c"], DEVICE_CONFIGS[form["c"]]["device_name"],
               "", form.get("how", "")]
        for field in ("Manufacturer", "Model", "S.N", "Status", "Location", "Date",
                      "S.N2", "Status2"):
            row += [form["land"].get(field, ""), (cert.get(field) or [""])[0]]
        problems = [f"{f}: {v[0]}" for f, v in form["judge"].items() if v[0] != "OK"]
        row += ["; ".join(problems) or "all OK", "", ""]
        sample.append(row)
        _link(sample.cell(sample.max_row, 4), form["p"])
        confirm.add(sample.cell(sample.max_row, len(header) - 1))
    sample.auto_filter.ref = sample.dimensions

    # Headline numbers, on a sheet of their own at the front.
    overview = workbook.create_sheet("Overview", 0)
    overview.column_dimensions["A"].width = 44
    overview.column_dimensions["B"].width = 14
    overview.column_dimensions["C"].width = 10
    overview.append(["Calist archive audit", datetime.now().strftime("%d-%m-%Y %H:%M")])
    overview.append([])
    overview.append(["Forms audited", len(judged)])
    overview.append(["Forms unreadable", len(unreadable)])
    overview.append(["Device lists and other non-forms skipped", len(not_forms)])
    overview.append(["Files with an unknown device code", len(unknown)])
    overview.append([])
    overview.append(["Field checks", "count", "share"])
    total = sum(grand.values()) or 1
    for kind, count in grand.most_common():
        overview.append([kind, count, f"{count / total:.1%}"])
    overview.append([])
    overview.append(["MAP ERROR — the form shows where the right value is; Calist reads another cell"])
    overview.append(["WRONG VALUE — cannot be right (a caption, a date as a serial). Recurring "
                     "across a device-year -> listed under Map errors; isolated -> Form errors"])
    overview.append(["MISSING — blank or N.A everywhere on the form: the engineer left it empty"])
    overview.append(["UNDETERMINED — sources disagree and nothing settles it"])
    overview.append(["EXPECTED — explained, not a fault (e.g. certificate names the probe)"])
    _save_outside_customers(workbook, path, root)


def _save_outside_customers(workbook, path: Path, root: str) -> None:
    """Save the report — and refuse, loudly, to write anywhere near the forms.

    The audit's whole promise is that it never changes a customer file. A bug
    once broke that: a loop variable shadowed the destination and the report
    was saved over a customer's workbook. Checking the one destination right
    before the one write is what makes that class of mistake impossible
    rather than merely unlikely.
    """
    target = Path(path).resolve()
    if target.suffix.lower() != ".xlsx" or not target.name.startswith("audit-"):
        raise RuntimeError(f"refusing to save the report as {target}")
    for year in YEARS:
        forms = (Path(root) / f"Customers {year}").resolve()
        if target == forms or forms in target.parents:
            raise RuntimeError(f"refusing to save the report inside {forms}")
    workbook.save(target)


# ──────────────────────────────────────────────────────────────────────────────
# Before / after
# ──────────────────────────────────────────────────────────────────────────────

def diff(before_path: str, after_path: str) -> int:
    """The proof that a fix did no harm: per form, per field, before vs after.

    A **regression** is a field that was OK before and is not now, or whose
    value changed while it was confirmed by the form. It must be zero.
    """
    def load(path):
        with gzip.open(path, "rt", encoding="utf-8") as handle:
            forms = {f["p"]: f for f in json.load(handle)["forms"]}
        # Re-judge with today's rules on both sides, so that only what the
        # reader lands can differ — never how the audit happens to judge it.
        for form in forms.values():
            if "land" in form:
                form["judge"] = judge(form)
        return forms
    before, after = load(before_path), load(after_path)
    kinds = Counter()
    regressions, improved, refused = [], Counter(), []
    for path, old in before.items():
        new = after.get(path)
        if new is None or "judge" not in old:
            continue
        if "notform" in new:
            # Now recognised as a device list: nothing to land, by design.
            # Listed by name below so each refusal is seen, not assumed.
            refused.append((old["c"], path))
            continue
        for field, (verdict, *_rest) in old["judge"].items():
            now = (new.get("judge") or {}).get(field, ["MISSING"])
            old_value = old["land"].get(field, "")
            new_value = (new.get("land") or {}).get(field, "")
            if verdict == "OK" and now[0] != "OK":
                kinds["REGRESSION: was OK, now " + now[0]] += 1
                regressions.append((old["c"], field, old_value, new_value, path))
            elif verdict == "OK" and not same(field, old_value, new_value):
                kinds["REGRESSION: OK value changed"] += 1
                regressions.append((old["c"], field, old_value, new_value, path))
            elif verdict != "OK" and now[0] == "OK":
                kinds[f"fixed ({verdict} -> OK)"] += 1
                improved[old["c"]] += 1
            elif verdict == now[0]:
                kinds["unchanged"] += 1
            else:
                kinds[f"{verdict} -> {now[0]}"] += 1
    newly = [p for p in after if p not in before]
    print(f"forms compared: {len(before):,}   newly readable: {len(newly):,}")
    for kind, count in kinds.most_common():
        print(f"  {count:>7,}  {kind}")
    print(f"\nnow refused as not a form: {len(refused)}")
    for code, path in refused:
        print(f"  {code:5} {os.path.basename(path)}")
    print(f"\nREGRESSIONS: {len(regressions)}")
    for code, field, old, new, path in regressions[:30]:
        print(f"  {code:5} {field:12} {old[:26]!r:30} -> {new[:26]!r}  {os.path.basename(path)}")
    print("\nfixed per device:", improved.most_common(20))
    return 1 if regressions else 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("root", nargs="?", default=r"D:\MedCal Pro")
    parser.add_argument("--workers", type=int, default=max(1, (os.cpu_count() or 4) - 2))
    parser.add_argument("--limit", type=int, help="audit a random N forms only")
    parser.add_argument("--out", help=f"report folder (default <root>/{REPORT_DIR})")
    parser.add_argument("--diff", nargs=2, metavar=("BEFORE", "AFTER"))
    parser.add_argument("--report", metavar="SNAPSHOT",
                        help="rebuild the workbook from a saved snapshot")
    args = parser.parse_args()
    if args.diff:
        return diff(*args.diff)
    if args.report:
        with gzip.open(args.report, "rt", encoding="utf-8") as handle:
            snap = json.load(handle)
        for form in snap["forms"]:
            if "land" in form:
                form["judge"] = judge(form)
        name = Path(args.report).name.replace(".json.gz", ".xlsx")
        out_dir = Path(args.out) if args.out else Path(args.report).parent
        target = out_dir / name
        write_report(snap["forms"], [tuple(u) for u in snap["unknown"]],
                     target, snap["root"])
        print(f"workbook  {target}")
        return 0
    run(args.root, args.workers, args.limit, args.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
